#!/usr/bin/env python3
"""Shared Answer-Key S1–S3 parsing and question-bank mapping.

Maps official sheet columns (e.g. ``1A01中文答案``) onto ``questionDatabase.js``
pools ``S1_CH`` / ``S1_EN`` / ``S2_*`` / ``S3_*``. Question numbers come from
the PNG GUID (``MCQ12``, ``MCQEng11``, ``MCEngQ.09``), not array order.
"""
from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
QDB = ROOT / "js/plugins/questionDatabase.js"
PICTURES = ROOT / "img/pictures/quiz"
DECISIONS_PATH = ROOT / "tools/answer_key_picture_decisions.json"
SHEET_LETTERS_PATH = ROOT / "tools/answer_key_s1s3_letters.json"

LETTERS = "ABCD"
POOLS = ("S1_CH", "S1_EN", "S2_CH", "S2_EN", "S3_CH", "S3_EN")
ASCII_PREFIX = {
    "S1_CH": "S1/CH",
    "S1_EN": "S1/EN",
    "S2_CH": "S2/CH",
    "S2_EN": "S2/EN",
    "S3_CH": "S3/CH",
    "S3_EN": "S3/EN",
}
SHEET_PREFIX = {
    "S1ANS01-12": "S1",
    "S2ANS01-12": "S2",
    "S3ANS01-12": "S3",
}
HDR_RE = re.compile(r"^([0-9A-Z]+)(中文|英文)答案")
Q_RE = re.compile(r"^Q(\d+)$", re.I)


def qnum_from_guid(guid: str) -> int | None:
    name = Path(guid).name
    for pat in (
        r"MCQEng(\d+)",
        r"MCEngQ\.?(\d+)",
        r"MCQ(\d+)",
        r"(?:^|[^0-9])Q(\d+)",
    ):
        m = re.search(pat, name, re.I)
        if m:
            return int(m.group(1))
    return None


def chapter_from_note(pool: str, note: str) -> str:
    if pool.endswith("_EN") and str(note).endswith("EN"):
        return note[:-2]
    return str(note)


def is_disabled(q: dict) -> bool:
    return all(str(q.get(k) or "") == "?" for k in ("C_A", "A2", "A3", "A4"))


def other_letters(correct: str) -> tuple[str, str, str]:
    rest = [x for x in LETTERS if x != correct]
    return rest[0], rest[1], rest[2]


def picture_path(pool: str, guid: str) -> Path:
    return PICTURES / ASCII_PREFIX[pool] / guid


def load_qdb(path: Path = QDB) -> dict:
    js = r"""
import fs from 'fs';
import vm from 'vm';
const code = fs.readFileSync(process.argv[1], 'utf8');
const sandbox = { questionDatabase: null };
vm.runInNewContext(code.replace(/^\s*var\s+questionDatabase\s*=\s*/, 'questionDatabase = '), sandbox);
const out = {};
for (const [k, arr] of Object.entries(sandbox.questionDatabase)) {
  if (!Array.isArray(arr)) continue;
  out[k] = arr.map(q => ({
    Note: q.Note, GUID: q.GUID, Q: q.Q, A: q.A,
    C_A: q.C_A, A2: q.A2, A3: q.A3, A4: q.A4, A5_Why: q.A5_Why
  }));
}
process.stdout.write(JSON.stringify(out));
"""
    r = subprocess.run(
        ["node", "--input-type=module", "-e", js, str(path)],
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(r.stdout)


def parse_xlsx_letters(xlsx: Path) -> dict[str, str]:
    """Return compact map ``POOL|CHAPTER|QN → A–D`` from Answer-Key-S1-6.xlsx."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out: dict[str, str] = {}
    for sname, prefix in SHEET_PREFIX.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        headers: list[tuple[int, str, str]] = []
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(1, c).value or "").strip()
            m = HDR_RE.match(h)
            if not m:
                continue
            ch = m.group(1)
            lang = "CH" if m.group(2) == "中文" else "EN"
            headers.append((c, f"{prefix}_{lang}", ch))
        seen: set[tuple[str, str, int]] = set()
        blocks: list[tuple[int, int, str, str]] = []
        for i, (c, pool, ch) in enumerate(headers):
            if (pool, ch, c) in seen:
                continue
            if (
                i + 1 < len(headers)
                and headers[i + 1][1] == pool
                and headers[i + 1][2] == ch
                and headers[i + 1][0] == c + 1
            ):
                blocks.append((c, c + 1, pool, ch))
                seen.add((pool, ch, c))
                seen.add((pool, ch, c + 1))
        for r in range(2, ws.max_row + 1):
            for qcol, acol, pool, ch in blocks:
                qcell = str(ws.cell(r, qcol).value or "").strip()
                mq = Q_RE.match(qcell)
                if not mq:
                    continue
                ans = str(ws.cell(r, acol).value or "").strip().upper()
                if ans in LETTERS:
                    out[f"{pool}|{ch}|{int(mq.group(1))}"] = ans
    return out


def load_sheet_letters(xlsx: Path | None = None) -> dict[str, str]:
    if xlsx and xlsx.exists():
        return parse_xlsx_letters(xlsx)
    if SHEET_LETTERS_PATH.exists():
        return json.loads(SHEET_LETTERS_PATH.read_text(encoding="utf-8"))
    return {}


def load_decisions(path: Path = DECISIONS_PATH) -> dict[str, dict]:
    if not path.exists():
        return {}
    rows = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(rows, dict) and "decisions" in rows:
        rows = rows["decisions"]
    out = {}
    for row in rows:
        guid = row.get("guid")
        if guid:
            out[guid] = row
    return out


def iter_mapped_questions(db: dict, sheet: dict[str, str]) -> list[dict]:
    rows = []
    for pool in POOLS:
        for q in db.get(pool, []):
            note = q.get("Note") or ""
            ch = chapter_from_note(pool, note)
            qn = qnum_from_guid(q.get("GUID") or "")
            guid = q.get("GUID") or ""
            db_ans = str(q.get("C_A") or "").strip().upper()
            disabled = is_disabled(q)
            key = f"{pool}|{ch}|{qn}" if qn is not None else ""
            rec = {
                "pool": pool,
                "chapter": ch,
                "note": note,
                "qn": qn,
                "guid": guid,
                "db": db_ans,
                "sheet": sheet.get(key),
                "disabled": disabled,
                "why": q.get("A5_Why") or "",
                "pic": picture_path(pool, guid) if pool in ASCII_PREFIX and guid else None,
                "has_pic": False,
            }
            if rec["pic"] is not None:
                rec["has_pic"] = rec["pic"].is_file()
                rec["pic"] = str(rec["pic"])
            rows.append(rec)
    return rows


def classify(rows: list[dict], decisions: dict[str, dict]) -> dict:
    counts = {
        "bank_total": len(rows),
        "matched": 0,
        "mismatch_live": 0,
        "fixed": 0,
        "still_disabled": 0,
        "sheet_vs_pic_conflicts": 0,
        "no_sheet": 0,
        "guid_parse_fail": 0,
        "pic_missing": 0,
        "keep_db": 0,
        "reviewed": 0,
    }
    samples = {
        "fixed": [],
        "still_disabled": [],
        "sheet_vs_pic_conflicts": [],
        "mismatch_unreviewed": [],
    }
    for rec in rows:
        if rec["qn"] is None:
            counts["guid_parse_fail"] += 1
        if rec.get("pic") and not rec.get("has_pic"):
            counts["pic_missing"] += 1
        if rec["sheet"] is None:
            counts["no_sheet"] += 1
        dec = decisions.get(rec["guid"])
        resolved = bool(dec) and dec.get("action") in ("fix", "keep_db", "keep_disabled")
        if rec["disabled"]:
            counts["still_disabled"] += 1
        elif rec["sheet"] and rec["db"] == rec["sheet"]:
            counts["matched"] += 1
        elif rec["sheet"] and rec["db"] != rec["sheet"]:
            if not resolved:
                counts["mismatch_live"] += 1

        if not dec:
            if rec["sheet"] and not rec["disabled"] and rec["db"] != rec["sheet"]:
                if len(samples["mismatch_unreviewed"]) < 30:
                    samples["mismatch_unreviewed"].append(rec)
            continue
        counts["reviewed"] += 1
        action = dec.get("action")
        pic_ans = dec.get("picture_answer")
        if isinstance(pic_ans, str):
            pic_ans = pic_ans.strip().upper() or None
        if action == "fix":
            counts["fixed"] += 1
            if len(samples["fixed"]) < 40:
                samples["fixed"].append({**rec, "decision": dec})
        elif action == "keep_db":
            counts["keep_db"] += 1
        if pic_ans and rec["sheet"] and pic_ans != rec["sheet"]:
            counts["sheet_vs_pic_conflicts"] += 1
            if len(samples["sheet_vs_pic_conflicts"]) < 40:
                samples["sheet_vs_pic_conflicts"].append({**rec, "decision": dec})
        if action == "keep_disabled" and len(samples["still_disabled"]) < 20:
            samples["still_disabled"].append({**rec, "decision": dec})
    return {"counts": counts, "samples": samples}

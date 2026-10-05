#!/usr/bin/env python3
"""Verify school-content fixes: S2 sample IDs, 2A05 disabled, red/yellow dialogue.

Usage:
  python3 tools/verify_school_content_fixes.py
  python3 tools/verify_school_content_fixes.py --xlsx /path/to/Answer-Key-S1-6.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
QDB = ROOT / "js/plugins/questionDatabase.js"
CE = ROOT / "data/CommonEvents.json"
MAP094 = ROOT / "data/Map094.json"
AUTH = ROOT / "tools/AUTH.md"

SAMPLE_CHECKS = [
    {
        "id": "2A05 MCQ1",
        "guid": "2A05/JSMATH2A05MCQ1.png",
        "expect": "disabled",
        "why": "題圖正解為 A(1,5)，官方/舊庫標 B(2,3)；全章暫停用",
    },
    {
        "id": "2A05 MCQ2",
        "guid": "2A05/JSMATH2A05MCQ2.png",
        "expect": "disabled",
        "why": "四個選項皆不滿足聯立方程",
    },
    {
        "id": "2A05 MCQ3",
        "guid": "2A05/JSMATH2A05MCQ3.png",
        "expect": "disabled",
        "why": "四個選項皆不滿足聯立方程",
    },
    {
        "id": "2A05 MCQ21",
        "guid": "2A05/JSMATH2A05MCQ21.png",
        "expect": "disabled",
        "why": "蘋果橙價錢正解 $41 不在 A–D（舊標 C $39）",
    },
    {
        "id": "2A05 MCQ51",
        "guid": "2A05/JSMATH2A05MCQ51.png",
        "expect": "disabled",
        "why": "正解約 (3.2, 3.6) 不在 A–D",
    },
    {
        "id": "2A05 MCQEng11",
        "guid": "2A05EN/JSMATH2A05MCQEng11.png",
        "expect": "disabled",
        "why": "英文聯立方程四選皆錯",
    },
    {
        "id": "2A02 MCQ6",
        "guid": "2A02/JSMATH2A02MCQ6.png",
        "expect": "D",
        "why": "最高次為 D 的 x^3 y^5（8 次），不是 B（6 次）",
    },
    {
        "id": "2A03 MCQ51",
        "guid": "2A03/JSMATH2A03MCQ51.png",
        "expect": "C",
        "why": "14x-28 = 14(x-2)，不是 A 的 7(x-4)",
    },
]

FORBIDDEN_PHRASES = [
    "彼母之劇情殺",
    "我D)@!$*@#$_L@@#$(+!(*LQ&Y$M!Q@)C(",
    "你他($%&@#(#*&$y*q",
    "形同廢人",
    "糞game",
    "數學腦殘粉",
    "數學白痴",
]

REQUIRED_PHRASES = [
    "呢個劇情殺都太犯規啦…",
    "可、可惡啊——！",
    "呢個陷阱都太狠啦……",
    "生活都亂晒套！",
    "爛遊戲",
    "數學死忠粉",
    "數學苦手",
]


def load_qdb_min() -> dict:
    import subprocess

    js = r"""
import fs from 'fs';
import vm from 'vm';
const code = fs.readFileSync(process.argv[1], 'utf8');
const sandbox = { questionDatabase: null };
vm.runInNewContext(code.replace(/^\s*var\s+questionDatabase\s*=\s*/, 'questionDatabase = '), sandbox);
const out = {};
for (const [k, arr] of Object.entries(sandbox.questionDatabase)) {
  if (!Array.isArray(arr)) continue;
  out[k] = arr.map(q => ({Note:q.Note, GUID:q.GUID, C_A:q.C_A, A2:q.A2, A3:q.A3, A4:q.A4, A5_Why:q.A5_Why}));
}
process.stdout.write(JSON.stringify(out));
"""
    r = subprocess.run(
        ["node", "--input-type=module", "-e", js, str(QDB)],
        check=True,
        capture_output=True,
        text=True,
    )
    return json.loads(r.stdout)


def find_q(db: dict, guid: str):
    for arr in db.values():
        for q in arr:
            if q.get("GUID") == guid:
                return q
    return None


def is_disabled(q: dict) -> bool:
    return all(str(q.get(k) or "") == "?" for k in ("C_A", "A2", "A3", "A4"))


def check_samples(db: dict) -> list[str]:
    errors = []
    for row in SAMPLE_CHECKS:
        q = find_q(db, row["guid"])
        if not q:
            errors.append(f"{row['id']}: 找不到 GUID {row['guid']}")
            continue
        if row["expect"] == "disabled":
            if not is_disabled(q):
                errors.append(f"{row['id']}: 應停用，實際 C_A={q.get('C_A')}")
        elif str(q.get("C_A")) != row["expect"]:
            errors.append(f"{row['id']}: 應為 {row['expect']}，實際 C_A={q.get('C_A')}")
    return errors


def check_chapter_disabled(db: dict) -> list[str]:
    errors = []
    for cat, notes in (("S2_CH", "2A05"), ("S2_EN", "2A05EN")):
        qs = [q for q in db[cat] if q.get("Note") == notes]
        if len(qs) != 100:
            errors.append(f"{cat} {notes} 預期 100 題，實際 {len(qs)}")
        live = [q for q in qs if not is_disabled(q)]
        if live:
            errors.append(f"{cat} {notes} 仍有 {len(live)} 題未停用")
    s1_live_2a05 = [
        q
        for q in db.get("S1_CH", []) + db.get("S1_EN", [])
        if str(q.get("Note", "")).startswith("2A05")
    ]
    if s1_live_2a05:
        errors.append("S1 不應出現 2A05")
    return errors


def check_s1_not_question_marked(db: dict) -> list[str]:
    errors = []
    for cat in ("S1_CH", "S1_EN"):
        bad = [q for q in db[cat] if is_disabled(q) or q.get("C_A") == "?"]
        if bad:
            errors.append(f"{cat} 不應被這次修正停用，卻有 {len(bad)} 題 C_A=?")
    return errors


def check_dialogue() -> list[str]:
    errors = []
    blob = "\n".join(
        p.read_text(encoding="utf-8")
        for p in (CE, MAP094, ROOT / "data/Map001.json", ROOT / "data/Map035.json")
    )
    for phrase in FORBIDDEN_PHRASES:
        if phrase in blob:
            errors.append(f"仍殘留紅/黃線用語：{phrase}")
    for phrase in REQUIRED_PHRASES:
        if phrase not in blob:
            errors.append(f"缺少應出現的校園用語：{phrase}")
    ce = json.loads(CE.read_text(encoding="utf-8"))
    ce7 = next(ev for ev in ce if ev and ev.get("id") == 7)
    # PR #5 not merged: do not rewrite CE7 flow in this PR.
    if ce7.get("name") != "地道劇情1":
        errors.append("CE7 名稱被改動")
    return errors


def check_auth_untouched() -> list[str]:
    import subprocess

    r = subprocess.run(
        ["git", "diff", "--name-only", "HEAD", "--", "tools/AUTH.md"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=True,
    )
    if r.stdout.strip():
        return ["不應改動 tools/AUTH.md"]
    if not AUTH.exists():
        return ["找不到 tools/AUTH.md"]
    return []


def optional_xlsx_note(xlsx: Path | None) -> str:
    if not xlsx or not xlsx.exists():
        return "未提供官方 xlsx（略過字母對照備註）"
    try:
        import openpyxl
    except ImportError:
        return "未安裝 openpyxl（略過字母對照備註）"
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["S2ANS01-12"]
    # 2A05 Q1 中文 = col 18 (R1 headers 17/18)
    q1 = str(ws.cell(2, 18).value or "").strip().upper()
    q6_2a02 = str(ws.cell(7, 6).value or "").strip().upper()
    return (
        f"官方表 2A05 Q1 中文={q1}（題圖數學為 A，故不跟表）；"
        f" 2A02 Q6 中文={q6_2a02}（題圖數學為 D，故不跟表）"
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=None)
    args = ap.parse_args()

    db = load_qdb_min()
    errors = []
    errors += check_samples(db)
    errors += check_chapter_disabled(db)
    errors += check_s1_not_question_marked(db)
    errors += check_dialogue()
    errors += check_auth_untouched()

    print("== 2A05 / 已知樣本檢查 ==")
    for row in SAMPLE_CHECKS:
        q = find_q(db, row["guid"])
        status = "MISSING"
        if q:
            status = "DISABLED" if is_disabled(q) else f"C_A={q.get('C_A')}"
        print(f"- {row['id']}: {status}  ({row['why']})")
    print()
    print(optional_xlsx_note(args.xlsx))
    print()
    if errors:
        print("FAILED:")
        for e in errors:
            print(" -", e)
        return 1
    print("OK: 樣本、2A05 停用、S1 未停用、紅黃線、AUTH 均通過")
    return 0


if __name__ == "__main__":
    sys.exit(main())
